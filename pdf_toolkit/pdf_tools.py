import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from docx import Document
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from xml.sax.saxutils import escape
from docx.text.paragraph import Paragraph as DocxParagraph
from pypdf import PdfWriter
from docx.table import Table as DocxTable
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from pypdf import PdfReader, PdfWriter
from PIL import Image, ImageOps
import warnings
Image.MAX_IMAGE_PIXELS = 40_000_000
warnings.simplefilter("error", Image.DecompressionBombWarning)
import cv2
import numpy as np
import pymupdf as fitz
import ocrmypdf
import img2pdf
from PIL import Image
from pptx import Presentation
from pptx.util import Emu
import pdfplumber
from openpyxl import Workbook
import ocrmypdf
from pyhanko.sign import signers
from pyhanko.pdf_utils.incremental_writer import IncrementalPdfFileWriter
import pymupdf as fitz
import difflib
import pymupdf as fitz
import pymupdf as fitz
import anthropic
import pymupdf as fitz
import ast


def _parse_pages(spec, total):
    """'1-3,5,8-' -> [0,1,2,4,7,...] (zero-based). Also accepts a list of ints (1-based)."""
    if isinstance(spec, (list, tuple)):
        return [int(p) - 1 for p in spec if 1 <= int(p) <= total]
    pages = []
    for part in str(spec).replace(" ", "").split(","):
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            start = int(a) if a else 1
            end = int(b) if b else total
            pages.extend(range(start - 1, min(end, total)))
        else:
            n = int(part)
            if 1 <= n <= total:
                pages.append(n - 1)
    return pages


def _docx_to_pdf_python(input_path, output_path):
    """Basic DOCX-to-PDF conversion using python-docx + ReportLab (no LibreOffice).

    This preserves readable text, paragraph order, headings, and simple tables.
    It is not a pixel-perfect Office renderer and may omit images, complex layouts,
    headers/footers, tracked changes, and some formatting.
    """

    docx = Document(str(input_path))
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="DocBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=10, leading=14, spaceAfter=5))
    styles.add(ParagraphStyle(name="DocHeading1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=18, leading=22, spaceBefore=9, spaceAfter=7))
    styles.add(ParagraphStyle(name="DocHeading2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=14, leading=18, spaceBefore=7, spaceAfter=5))
    styles.add(ParagraphStyle(name="DocHeading3", parent=styles["Heading3"], fontName="Helvetica-Bold", fontSize=12, leading=15, spaceBefore=5, spaceAfter=4))
    story = []

    body = docx.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            p = DocxParagraph(child, docx)
            raw = p.text or ""
            if not raw.strip():
                story.append(Spacer(1, 4))
                continue
            txt = escape(raw).replace("\n", "<br/>")
            style_name = (p.style.name if p.style else "")
            if style_name.startswith("Heading 1") or style_name == "Title":
                st = styles["DocHeading1"]
            elif style_name.startswith("Heading 2"):
                st = styles["DocHeading2"]
            elif style_name.startswith("Heading 3"):
                st = styles["DocHeading3"]
            else:
                st = styles["DocBody"]
            if p.alignment is not None:
                align = {1: TA_CENTER, 2: TA_RIGHT, 0: TA_LEFT}.get(int(p.alignment), None)
                if align is not None and st.name == "DocBody":
                    st = ParagraphStyle("align_" + str(align), parent=st, alignment=align)
            story.append(Paragraph(txt, st))
        elif isinstance(child, CT_Tbl):
            tbl = DocxTable(child, docx)
            data = []
            for row in tbl.rows:
                cells = []
                for cell in row.cells:
                    cell_text = escape(cell.text or "").replace("\n", "<br/>")
                    cells.append(Paragraph(cell_text or " ", styles["DocBody"]))
                data.append(cells)
            if data:
                table = Table(data, repeatRows=1, hAlign="LEFT")
                table.setStyle(TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]))
                story.extend([Spacer(1, 5), table, Spacer(1, 8)])
    if not story:
        story = [Paragraph("This document contains no extractable text.", styles["DocBody"])]
    pdf = SimpleDocTemplate(str(output_path), pagesize=A4,
                            rightMargin=18*mm, leftMargin=18*mm,
                            topMargin=18*mm, bottomMargin=18*mm,
                            title=Path(input_path).stem)
    pdf.build(story)
    return output_path


def _libreoffice_convert(input_path, output_path):
    """Use LibreOffice when available; otherwise offer a basic DOCX-only fallback."""
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if exe:
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run(
                [exe, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(input_path)],
                check=True, capture_output=True,
            )
            produced = Path(tmp) / (Path(input_path).stem + ".pdf")
            if not produced.is_file():
                raise RuntimeError("LibreOffice did not create the expected PDF.")
            shutil.move(str(produced), output_path)
        return output_path

    ext = Path(input_path).suffix.lower()
    if ext == ".docx":
        return _docx_to_pdf_python(input_path, output_path)
    raise RuntimeError(
        f"LibreOffice is not installed. Built-in no-LibreOffice conversion currently supports .docx only; "
        f"this file is {ext or 'unknown format'}. Convert it to .docx first, or install LibreOffice for full Office-format support."
    )

def merge_pdf(inputs, output):
    """Merge several PDFs into one."""
    writer = PdfWriter()
    for f in inputs:
        writer.append(f)
    with open(output, "wb") as fh:
        writer.write(fh)
    return output


def split_pdf(input_path, out_dir, ranges=None):
    """Split a PDF. ranges=None -> one file per page.
    ranges='1-3,4-6' -> one file per comma-separated range."""
    reader = PdfReader(input_path)
    total = len(reader.pages)
    os.makedirs(out_dir, exist_ok=True)
    stem = Path(input_path).stem
    groups = [[i] for i in range(total)] if not ranges else [
        _parse_pages(r, total) for r in str(ranges).split(",")
    ]
    outputs = []
    for n, group in enumerate(groups, 1):
        w = PdfWriter()
        for i in group:
            w.add_page(reader.pages[i])
        path = os.path.join(out_dir, f"{stem}_part{n}.pdf")
        with open(path, "wb") as fh:
            w.write(fh)
        outputs.append(path)
    return outputs


def remove_pages(input_path, output, pages):
    """Remove pages, e.g. pages='2,4-6'."""
    reader = PdfReader(input_path)
    drop = set(_parse_pages(pages, len(reader.pages)))
    w = PdfWriter()
    for i, p in enumerate(reader.pages):
        if i not in drop:
            w.add_page(p)
    with open(output, "wb") as fh:
        w.write(fh)
    return output


def extract_pages(input_path, output, pages):
    """Keep only the given pages, e.g. pages='1,3-5'."""
    reader = PdfReader(input_path)
    w = PdfWriter()
    for i in _parse_pages(pages, len(reader.pages)):
        w.add_page(reader.pages[i])
    with open(output, "wb") as fh:
        w.write(fh)
    return output


def organize_pdf(input_path, output, order):
    """Reorder (and optionally duplicate/drop) pages. order=[3,1,2] (1-based)."""
    reader = PdfReader(input_path)
    w = PdfWriter()
    for i in _parse_pages(order, len(reader.pages)):
        w.add_page(reader.pages[i])
    with open(output, "wb") as fh:
        w.write(fh)
    return output


def scan_to_pdf(images, output, enhance=True):
    """Turn photos of documents into a clean, scanner-style PDF."""
    pages = []
    for path in images:
        img = Image.open(path)
        img = ImageOps.exif_transpose(img).convert("RGB")
        if enhance:
            try:
                gray = cv2.cvtColor(np.array(img), cv2.COLOR_RGB2GRAY)
                gray = cv2.adaptiveThreshold(
                    gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 31, 15
                )
                img = Image.fromarray(gray).convert("RGB")
            except ImportError:
                img = ImageOps.autocontrast(ImageOps.grayscale(img)).convert("RGB")
        pages.append(img)
    pages[0].save(output, save_all=True, append_images=pages[1:], resolution=200)
    return output


# ======================================================================
# OPTIMIZE PDF
# ======================================================================
def compress_pdf(input_path, output, quality="ebook"):
    """quality: screen (smallest) | ebook | printer | prepress.
    Uses Ghostscript if available, else PyMuPDF structural compression."""
    gs = shutil.which("gs") or shutil.which("gswin64c")
    if gs:
        subprocess.run(
            [gs, "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.4",
             f"-dPDFSETTINGS=/{quality}", "-dNOPAUSE", "-dQUIET", "-dBATCH",
             f"-sOutputFile={output}", input_path],
            check=True,
        )
    else:
        doc = fitz.open(input_path)
        doc.save(output, garbage=4, deflate=True, clean=True)
        doc.close()
    return output


def repair_pdf(input_path, output):
    """Rebuild a damaged PDF (PyMuPDF repairs xref on open)."""
    doc = fitz.open(input_path)
    doc.save(output, garbage=4, clean=True, deflate=True)
    doc.close()
    return output


def ocr_pdf(input_path, output, language="eng"):
    """Make a scanned PDF searchable (needs Tesseract)."""
    ocrmypdf.ocr(input_path, output, language=language, force_ocr=True, deskew=True)
    return output


# ======================================================================
# CONVERT TO PDF
# ======================================================================
def jpg_to_pdf(images, output):
    """One or more images -> one PDF."""
    try:
        with open(output, "wb") as fh:
            fh.write(img2pdf.convert([str(i) for i in images]))
    except Exception:
        imgs = [Image.open(i).convert("RGB") for i in images]
        imgs[0].save(output, save_all=True, append_images=imgs[1:])
    return output


def word_to_pdf(input_path, output):
    return _libreoffice_convert(input_path, output)


def powerpoint_to_pdf(input_path, output):
    return _libreoffice_convert(input_path, output)


def excel_to_pdf(input_path, output):
    return _libreoffice_convert(input_path, output)


def _deny_external_resources(url):
    """WeasyPrint fetcher: block network and local-file resource access."""
    from urllib.parse import urlsplit
    from weasyprint.urls import URLFetchingError
    parsed = urlsplit(url)
    if parsed.scheme.lower() != "data":
        raise URLFetchingError("External and local resources are disabled during HTML-to-PDF conversion.")
    from weasyprint import default_url_fetcher
    return default_url_fetcher(url)


def html_to_pdf(source, output):
    """Convert a user-uploaded HTML file without network or local-file fetching."""
    from weasyprint import HTML
    path = Path(str(source))
    if not path.is_file() or path.suffix.lower() not in {".html", ".htm"}:
        raise ValueError("Upload a valid .html or .htm file.")
    HTML(filename=str(path), url_fetcher=_deny_external_resources).write_pdf(output)
    return output


# ======================================================================
# CONVERT FROM PDF
# ======================================================================
def pdf_to_jpg(input_path, out_dir, dpi=150, quality=90):
    os.makedirs(out_dir, exist_ok=True)
    doc = fitz.open(input_path)
    outputs = []
    for i, page in enumerate(doc, 1):
        pix = page.get_pixmap(dpi=dpi)
        path = os.path.join(out_dir, f"page_{i:03d}.jpg")
        pix.save(path, jpg_quality=quality)
        outputs.append(path)
    doc.close()
    return outputs


def pdf_to_word(input_path, output):
    from pdf2docx import Converter
    cv = Converter(input_path)
    cv.convert(output)
    cv.close()
    return output


def pdf_to_powerpoint(input_path, output, dpi=150):
    """Each PDF page becomes a full-slide image."""
    doc = fitz.open(input_path)
    prs = Presentation()
    first = doc[0].rect
    prs.slide_width = Emu(int(first.width * 12700))
    prs.slide_height = Emu(int(first.height * 12700))
    blank = prs.slide_layouts[6]
    with tempfile.TemporaryDirectory() as tmp:
        for i, page in enumerate(doc):
            img = os.path.join(tmp, f"{i}.png")
            page.get_pixmap(dpi=dpi).save(img)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(img, 0, 0, prs.slide_width, prs.slide_height)
    prs.save(output)
    doc.close()
    return output


def pdf_to_excel(input_path, output):
    """Extract tables (one sheet per table)."""
    wb = Workbook()
    wb.remove(wb.active)
    count = 0
    with pdfplumber.open(input_path) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            for tno, table in enumerate(page.extract_tables(), 1):
                count += 1
                ws = wb.create_sheet(f"P{pno}_T{tno}"[:31])
                for row in table:
                    ws.append([c if c is not None else "" for c in row])
    if count == 0:
        ws = wb.create_sheet("No tables found")
        ws["A1"] = "No tables detected in this PDF."
    wb.save(output)
    return output


def pdf_to_pdfa(input_path, output):
    """Convert to PDF/A (archival)."""
    try:
        ocrmypdf.ocr(input_path, output, output_type="pdfa", skip_text=True)
        return output
    except Exception:
        gs = shutil.which("gs") or shutil.which("gswin64c")
        if not gs:
            raise RuntimeError("Need ocrmypdf or Ghostscript for PDF/A conversion.")
        subprocess.run(
            [gs, "-dPDFA=2", "-dBATCH", "-dNOPAUSE", "-dQUIET",
             "-sColorConversionStrategy=RGB", "-sProcessColorModel=DeviceRGB",
             "-dPDFACompatibilityPolicy=1", "-sDEVICE=pdfwrite",
             f"-sOutputFile={output}", input_path],
            check=True,
        )
        return output


# ======================================================================
# EDIT PDF
# ======================================================================
def rotate_pdf(input_path, output, angle=90, pages=None):
    """angle: 90, 180, 270. pages=None -> all pages."""
    reader = PdfReader(input_path)
    targets = set(_parse_pages(pages, len(reader.pages))) if pages else set(range(len(reader.pages)))
    w = PdfWriter()
    for i, p in enumerate(reader.pages):
        if i in targets:
            p.rotate(int(angle))
        w.add_page(p)
    with open(output, "wb") as fh:
        w.write(fh)
    return output


def add_page_numbers(input_path, output, position="bottom-center",
                     fmt="Page {n} of {total}", start=1, fontsize=10):
    """position: bottom-center | bottom-right | bottom-left | top-center | top-right | top-left"""
    doc = fitz.open(input_path)
    total = len(doc)
    for i, page in enumerate(doc):
        text = fmt.format(n=i + start, total=total + start - 1)
        r = page.rect
        width = fitz.get_text_length(text, fontsize=fontsize)
        vert, horiz = position.split("-")
        y = r.height - 25 if vert == "bottom" else 30
        x = {"left": 36, "center": (r.width - width) / 2, "right": r.width - width - 36}[horiz]
        page.insert_text((x, y), text, fontsize=fontsize, color=(0, 0, 0))
    doc.save(output)
    doc.close()
    return output


def add_watermark(input_path, output, text="CONFIDENTIAL", fontsize=60,
                  color=(0.6, 0.6, 0.6), opacity=0.3, angle=45, image=None):
    """Text watermark (diagonal) or image watermark (image='logo.png')."""
    doc = fitz.open(input_path)
    for page in doc:
        r = page.rect
        if image:
            w, h = r.width * 0.5, r.height * 0.5
            rect = fitz.Rect((r.width - w) / 2, (r.height - h) / 2,
                             (r.width + w) / 2, (r.height + h) / 2)
            page.insert_image(rect, filename=image, keep_proportion=True, overlay=True)
        else:
            length = fitz.get_text_length(text, fontsize=fontsize)
            point = fitz.Point((r.width - length) / 2, r.height / 2)
            page.insert_text(
                point, text, fontsize=fontsize, color=color, fill_opacity=opacity,
                morph=(fitz.Point(r.width / 2, r.height / 2), fitz.Matrix(-angle)),
                overlay=True,
            )
    doc.save(output)
    doc.close()
    return output


def crop_pdf(input_path, output, left=0, top=0, right=0, bottom=0):
    """Crop margins (in PDF points, 72 = 1 inch)."""
    doc = fitz.open(input_path)
    for page in doc:
        r = page.rect
        page.set_cropbox(fitz.Rect(r.x0 + left, r.y0 + top, r.x1 - right, r.y1 - bottom))
    doc.save(output)
    doc.close()
    return output


def edit_pdf(input_path, output, replacements=None, add_text=None, add_image=None):
    """
    replacements : {"old text": "new text"}
    add_text     : [{"page": 1, "x": 72, "y": 72, "text": "Hello", "size": 12}]
    add_image    : [{"page": 1, "rect": [50, 50, 200, 150], "path": "logo.png"}]
    """
    doc = fitz.open(input_path)
    for old, new in (replacements or {}).items():
        for page in doc:
            for rect in page.search_for(old):
                page.add_redact_annot(rect, text=new, fontsize=11, fill=(1, 1, 1))
            page.apply_redactions()
    for item in add_text or []:
        page = doc[item.get("page", 1) - 1]
        page.insert_text((item["x"], item["y"]), item["text"],
                         fontsize=item.get("size", 12), color=item.get("color", (0, 0, 0)))
    for item in add_image or []:
        page = doc[item.get("page", 1) - 1]
        page.insert_image(fitz.Rect(*item["rect"]), filename=item["path"])
    doc.save(output)
    doc.close()
    return output


def list_form_fields(input_path):
    """Return {field_name: (type, current_value)} for a fillable PDF."""
    doc = fitz.open(input_path)
    fields = {}
    for page in doc:
        for w in page.widgets() or []:
            fields[w.field_name] = (w.field_type_string, w.field_value)
    doc.close()
    return fields


def fill_form(input_path, output, values):
    """values = {"Name": "John", "Agree": True}"""
    doc = fitz.open(input_path)
    for page in doc:
        for w in page.widgets() or []:
            if w.field_name in values:
                w.field_value = values[w.field_name]
                w.update()
    doc.save(output)
    doc.close()
    return output


def add_form_field(input_path, output, page=1, name="field1", kind="text",
                   rect=(72, 72, 272, 100)):
    """kind: text | checkbox"""
    doc = fitz.open(input_path)
    pg = doc[page - 1]
    w = fitz.Widget()
    w.rect = fitz.Rect(*rect)
    w.field_name = name
    w.field_type = fitz.PDF_WIDGET_TYPE_CHECKBOX if kind == "checkbox" else fitz.PDF_WIDGET_TYPE_TEXT
    w.border_color = (0, 0, 0)
    pg.add_widget(w)
    doc.save(output)
    doc.close()
    return output


# ======================================================================
# PDF SECURITY
# ======================================================================
def unlock_pdf(input_path, output, password):
    reader = PdfReader(input_path)
    if reader.is_encrypted and not reader.decrypt(password):
        raise ValueError("Wrong password.")
    w = PdfWriter()
    for p in reader.pages:
        w.add_page(p)
    with open(output, "wb") as fh:
        w.write(fh)
    return output


def protect_pdf(input_path, output, user_password, owner_password=None):
    reader = PdfReader(input_path)
    w = PdfWriter()
    for p in reader.pages:
        w.add_page(p)
    w.encrypt(user_password=user_password,
              owner_password=owner_password or user_password,
              algorithm="AES-256")
    with open(output, "wb") as fh:
        w.write(fh)
    return output


def sign_pdf(input_path, output, signature_image, page=1, rect=(350, 700, 550, 780)):
    """Visible signature: stamps a signature image on a page."""
    doc = fitz.open(input_path)
    doc[page - 1].insert_image(fitz.Rect(*rect), filename=signature_image)
    doc.save(output)
    doc.close()
    return output


def sign_pdf_digital(input_path, output, p12_file, p12_password, field_name="Signature1"):
    """Cryptographic (PKCS#12) digital signature using pyHanko."""
    signer = signers.SimpleSigner.load_pkcs12(p12_file, passphrase=p12_password.encode())
    with open(input_path, "rb") as inf:
        w = IncrementalPdfFileWriter(inf)
        with open(output, "wb") as outf:
            signers.sign_pdf(w, signers.PdfSignatureMetadata(field_name=field_name),
                             signer=signer, output=outf)
    return output


def redact_pdf(input_path, output, terms):
    """Permanently remove text. terms = ["John Doe", "123-45-6789"]"""
    doc = fitz.open(input_path)
    for page in doc:
        for term in terms:
            for rect in page.search_for(term):
                page.add_redact_annot(rect, fill=(0, 0, 0))
        page.apply_redactions()
    doc.save(output, garbage=4, deflate=True)
    doc.close()
    return output


def compare_pdf(pdf_a, pdf_b, output_html="comparison.html"):
    """Side-by-side text diff of two PDFs -> HTML report."""

    def lines(path):
        d = fitz.open(path)
        out = []
        for i, p in enumerate(d, 1):
            out.append(f"--- Page {i} ---")
            out.extend(p.get_text().splitlines())
        d.close()
        return out

    html = difflib.HtmlDiff(wrapcolumn=80).make_file(
        lines(pdf_a), lines(pdf_b), Path(pdf_a).name, Path(pdf_b).name
    )
    Path(output_html).write_text(html, encoding="utf-8")
    return output_html


# ======================================================================
# PDF INTELLIGENCE
# ======================================================================
def _pdf_text(path):
    doc = fitz.open(path)
    pages = [p.get_text() for p in doc]
    doc.close()
    return pages


def _ask_claude(prompt, max_tokens=4096):
    client = anthropic.Anthropic()  # needs ANTHROPIC_API_KEY env var
    msg = client.messages.create(
        model=os.environ.get("CLAUDE_MODEL", "claude-sonnet-5-5"),
        max_tokens=max_tokens,
        messages=[{"role": "user", "content": prompt}],
    )
    return msg.content[0].text


def ai_summarizer(input_path, length="medium"):
    """length: short | medium | detailed. Returns summary text."""
    text = "\n\n".join(_pdf_text(input_path))[:150000]
    return _ask_claude(
        f"Summarize this document ({length} length). Use clear headings and bullets "
        f"where helpful.\n\n<document>\n{text}\n</document>"
    )


def translate_pdf(input_path, output, target_language="Hindi"):
    """Translate page by page into a new PDF."""
    src = _pdf_text(input_path)
    out = fitz.open()
    for text in src:
        if not text.strip():
            out.new_page()
            continue
        translated = _ask_claude(
            f"Translate the following text into {target_language}. Keep the line/paragraph "
            f"structure. Return only the translation.\n\n{text}"
        )
        page = out.new_page()
        html = "<div style='font-size:11pt'>" + escape(translated).replace("\n", "<br>") + "</div>"
        page.insert_htmlbox(fitz.Rect(40, 40, page.rect.width - 40, page.rect.height - 40), html)
    out.save(output)
    out.close()
    return output


def pdf_to_markdown(input_path, output="output.md"):
    import pymupdf4llm
    md = pymupdf4llm.to_markdown(input_path)
    Path(output).write_text(md, encoding="utf-8")
    return output


TOOLS = {
    f.__name__: f for f in [
        merge_pdf, split_pdf, remove_pages, extract_pages, organize_pdf, scan_to_pdf,
        compress_pdf, repair_pdf, ocr_pdf,
        jpg_to_pdf, word_to_pdf, powerpoint_to_pdf, excel_to_pdf, html_to_pdf,
        pdf_to_jpg, pdf_to_word, pdf_to_powerpoint, pdf_to_excel, pdf_to_pdfa,
        rotate_pdf, add_page_numbers, add_watermark, crop_pdf, edit_pdf,
        list_form_fields, fill_form, add_form_field,
        unlock_pdf, protect_pdf, sign_pdf, sign_pdf_digital, redact_pdf, compare_pdf,
        ai_summarizer, translate_pdf, pdf_to_markdown,
    ]
}


def _coerce(arg):
    try:
        return ast.literal_eval(arg)
    except Exception:
        return arg


def main(argv):
    if len(argv) < 2 or argv[1] in ("-h", "--help", "--list"):
        print("Available tools:\n  " + "\n  ".join(sorted(TOOLS)))
        print("\nExample: python pdf_tools.py rotate_pdf in.pdf out.pdf 90")
        return
    name = argv[1]
    if name not in TOOLS:
        sys.exit(f"Unknown tool '{name}'. Use --list.")
    args, kwargs = [], {}
    for a in argv[2:]:
        if "=" in a and not a.startswith(("[", "{", "'", '"')) and a.split("=")[0].isidentifier():
            k, v = a.split("=", 1)
            kwargs[k] = _coerce(v)
        else:
            args.append(_coerce(a))
    result = TOOLS[name](*args, **kwargs)
    print("Done:", result)


if __name__ == "__main__":
    main(sys.argv)
